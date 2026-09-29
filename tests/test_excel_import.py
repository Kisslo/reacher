import importlib

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from reacher.cli import app
from reacher.db import connect, migrate
from reacher.excel.contract import COL, META_SHEET, SHEET
from reacher.excel.export import CallListRow, export_call_list

outcome_import = importlib.import_module("reacher.excel.import")


def _metadata(call_list_id: int) -> dict[str, str]:
    return {
        "call_list_id": str(call_list_id),
        "week": "2026w40",
        "seller": "Anna",
        "scoring_version": "test-v1",
        "generated_at": "2026-09-29T10:00:00+00:00",
    }


@pytest.fixture
def prepared(tmp_path):
    db_path = tmp_path / "outcomes.db"
    workbook_path = tmp_path / "returned.xlsx"
    conn = connect(db_path)
    migrate(conn)

    salon_ids = []
    for number in range(1, 5):
        salon_ids.append(
            conn.execute(
                "INSERT INTO salon (orgnr, name, first_seen_at, last_seen_at) "
                "VALUES (?, ?, '2026-09-29', '2026-09-29')",
                (f"556123456{number}", f"Salon {number}"),
            ).lastrowid
        )

    call_list_id = conn.execute(
        "INSERT INTO call_list (week, seller, scoring_version, created_at) "
        "VALUES ('2026w40', 'Anna', 'test-v1', '2026-09-29')"
    ).lastrowid
    row_ids = []
    for rank, salon_id in enumerate(salon_ids, start=1):
        row_ids.append(
            conn.execute(
                "INSERT INTO call_list_row "
                "(call_list_id, salon_id, rank, score, reasons, phone) "
                "VALUES (?, ?, ?, 1, '[]', ?)",
                (call_list_id, salon_id, rank, f"+461234567{rank}"),
            ).lastrowid
        )
    conn.commit()

    export_call_list(
        workbook_path,
        [
            CallListRow(row_id, rank, 1, f"Salon {rank}", "Stockholm", f"+461234567{rank}", "", "")
            for rank, row_id in enumerate(row_ids, start=1)
        ],
        _metadata(call_list_id),
    )
    yield conn, db_path, workbook_path, row_ids, call_list_id
    conn.close()


def set_cells(workbook_path, values):
    workbook = load_workbook(workbook_path)
    worksheet = workbook[SHEET]
    for row_number, (outcome, comment) in values.items():
        worksheet.cell(row=row_number + 1, column=COL["Utfall"]).value = outcome
        worksheet.cell(row=row_number + 1, column=COL["Kommentar"]).value = comment
    workbook.save(workbook_path)


def test_imports_outcomes_comments_and_suppressions(prepared):
    conn, _, workbook_path, row_ids, _ = prepared
    set_cells(
        workbook_path,
        {
            1: ("Intresserad", "Ringde tillbaka"),
            2: ("Registrerad", "Kund"),
            3: ("Spärra", "Vill inte bli kontaktad"),
        },
    )

    summary = outcome_import.import_outcomes(conn, workbook_path)

    assert summary.imported == 3
    assert summary.empty_outcomes == 1
    assert summary.outcomes == {"Intresserad": 1, "Registrerad": 1, "Spärra": 1}
    assert summary.suppressions_added == {"existing_customer": 1, "opt_out": 1}
    assert conn.execute("SELECT count(*) FROM outcome").fetchone()[0] == 3
    assert (
        conn.execute(
            "SELECT comment FROM outcome WHERE call_list_row_id = ?", (row_ids[0],)
        ).fetchone()[0]
        == "Ringde tillbaka"
    )


def test_unknown_outcomes_and_suspicious_comments_are_reported(prepared):
    conn, _, workbook_path, row_ids, _ = prepared
    set_cells(
        workbook_path,
        {
            1: ("Unknown", "ring inte"),
            2: ("Nej", "SPÄRRA enligt kunden"),
        },
    )

    summary = outcome_import.import_outcomes(conn, workbook_path)

    assert summary.unknown_outcomes == ("Unknown",)
    assert summary.manual_review_row_ids == tuple(row_ids[:2])
    assert summary.imported == 1
    assert conn.execute("SELECT count(*) FROM outcome").fetchone()[0] == 1


def test_unknown_row_id_is_reported_without_being_imported(prepared):
    conn, _, workbook_path, _, _ = prepared
    set_cells(workbook_path, {1: ("Nej", "")})
    workbook = load_workbook(workbook_path)
    workbook[SHEET].cell(row=2, column=COL["row_id"]).value = 9999
    workbook.save(workbook_path)

    summary = outcome_import.import_outcomes(conn, workbook_path)

    assert summary.unknown_row_ids == (9999,)
    assert summary.imported == 0
    assert conn.execute("SELECT count(*) FROM outcome").fetchone()[0] == 0


def test_reimport_is_idempotent_and_does_not_remove_opt_out(prepared):
    conn, _, workbook_path, row_ids, _ = prepared
    set_cells(workbook_path, {3: ("Spärra", "Block this salon")})

    first = outcome_import.import_outcomes(conn, workbook_path)
    second = outcome_import.import_outcomes(conn, workbook_path)
    assert first.suppressions_added == {"opt_out": 1}
    assert second.suppressions_added == {"opt_out": 0}

    set_cells(workbook_path, {3: ("Nej", "Changed later")})
    outcome_import.import_outcomes(conn, workbook_path)

    assert (
        conn.execute(
            "SELECT outcome FROM outcome WHERE call_list_row_id = ?", (row_ids[2],)
        ).fetchone()[0]
        == "Nej"
    )
    assert (
        conn.execute("SELECT count(*) FROM suppression WHERE reason = 'opt_out'").fetchone()[0] == 1
    )


def test_metadata_mismatch_is_rejected_before_writes(prepared):
    conn, _, workbook_path, _, _ = prepared
    set_cells(workbook_path, {1: ("Nej", "")})
    workbook = load_workbook(workbook_path)
    meta = workbook[META_SHEET]
    for row in range(2, meta.max_row + 1):
        if meta.cell(row=row, column=1).value == "week":
            meta.cell(row=row, column=2).value = "2026w41"
    workbook.save(workbook_path)

    with pytest.raises(outcome_import.OutcomeImportError):
        outcome_import.import_outcomes(conn, workbook_path)
    assert conn.execute("SELECT count(*) FROM outcome").fetchone()[0] == 0


def test_cli_imports_workbook_and_prints_summary(prepared):
    conn, db_path, workbook_path, _, _ = prepared
    set_cells(workbook_path, {1: ("Registrerad", "")})
    conn.close()

    result = CliRunner().invoke(
        app,
        ["import-outcomes", str(workbook_path), "--db", str(db_path)],
    )

    assert result.exit_code == 0
    assert "Imported outcomes: 1" in result.stdout
    assert "Registrerad: 1" in result.stdout