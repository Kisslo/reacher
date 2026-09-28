import pytest
from openpyxl import load_workbook

from reacher.excel.contract import COL, COLUMN_WIDTHS, COLUMNS, HIDDEN, META_KEYS, META_SHEET, SHEET
from reacher.excel.export import CallListRow, export_call_list


def valid_metadata():
    return {
        "call_list_id": "42",
        "week": "2026w40",
        "seller": "Anna",
        "scoring_version": "1",
        "generated_at": "2026-09-28T10:30:00+00:00",
    }


def test_export_round_trip(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    rows = [
        CallListRow(
            row_id=17,
            rank=1,
            score=2.0,
            salon="Klipp & Co",
            area="Stockholm",
            phone="+46701234567",
            source="https://example.com/contact",
            reasons="Nystartad salong; 1-4 anställda",
        )
    ]

    export_call_list(output_path, rows, valid_metadata())

    workbook = load_workbook(output_path)

    worksheet = workbook[SHEET]
    assert [cell.value for cell in worksheet[1]] == list(COLUMNS)
    for column_name in HIDDEN:
        column_letter = worksheet.cell(row=1, column=COL[column_name]).column_letter
        assert worksheet.column_dimensions[column_letter].hidden is True
    assert worksheet["A2"].value == 17

    meta_sheet = workbook["_meta"]
    assert meta_sheet.sheet_state == "hidden"

    exported_keys = {
        meta_sheet.cell(row=row_number, column=1).value
        for row_number in range(2, meta_sheet.max_row + 1)
    }
    assert exported_keys == set(META_KEYS)


def test_export_writes_all_row_values(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    row = CallListRow(
        row_id=17,
        rank=1,
        score=2.0,
        salon="Klipp & Co",
        area="Stockholm",
        phone="+46701234567",
        source="https://example.com/contact",
        reasons="Nystartad salong",
    )

    export_call_list(output_path, [row], valid_metadata())

    worksheet = load_workbook(output_path)[SHEET]

    assert worksheet["A2"].value == 17
    assert worksheet["B2"].value == 1
    assert worksheet["C2"].value == 2.0
    assert worksheet["D2"].value == "Klipp & Co"
    assert worksheet["E2"].value == "Stockholm"
    assert worksheet["F2"].value == "+46701234567"
    assert worksheet["G2"].value == "https://example.com/contact"
    assert worksheet["H2"].value == "Nystartad salong"
    assert worksheet["I2"].value is None
    assert worksheet["J2"].value is None


def test_export_sets_column_widths_and_wraps_row_text(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    export_call_list(
        output_path,
        [
            CallListRow(
                row_id=17,
                rank=1,
                score=2.0,
                salon="Klipp & Co",
                area="Stockholm",
                phone="+46701234567",
                source="",
                reasons="Nystartad salong",
            )
        ],
        valid_metadata(),
    )

    worksheet = load_workbook(output_path)[SHEET]

    for column_name in COLUMNS:
        column_letter = worksheet.cell(row=1, column=COL[column_name]).column_letter
        assert worksheet.column_dimensions[column_letter].width == COLUMN_WIDTHS[column_name]

    assert worksheet["H2"].alignment.wrap_text is True
    assert worksheet["H2"].alignment.vertical == "top"


def test_only_outcome_and_comment_are_editable(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    export_call_list(
        output_path,
        [
            CallListRow(
                row_id=17,
                rank=1,
                score=2.0,
                salon="Klipp & Co",
                area="Stockholm",
                phone="+46701234567",
                source="",
                reasons="Nystartad salong",
            )
        ],
        valid_metadata(),
    )

    worksheet = load_workbook(output_path)[SHEET]

    assert worksheet.protection.sheet is True
    assert worksheet["I2"].protection.locked is False
    assert worksheet["J2"].protection.locked is False
    assert worksheet["A2"].protection.locked is True
    assert worksheet["D2"].protection.locked is True
    assert worksheet["I1"].protection.locked is True


def test_outcome_dropdown_has_exact_values(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    export_call_list(
        output_path,
        [
            CallListRow(
                row_id=17,
                rank=1,
                score=2.0,
                salon="Klipp & Co",
                area="Stockholm",
                phone="+46701234567",
                source="",
                reasons="Nystartad salong",
            )
        ],
        valid_metadata(),
    )

    worksheet = load_workbook(output_path)[SHEET]
    validations = worksheet.data_validations.dataValidation

    assert len(validations) == 1
    validation = validations[0]
    assert validation.type == "list"
    assert validation.allow_blank is True
    assert validation.formula1 == ('"Ej nådd,Intresserad,Registrerad,Nej,Spärra"')
    assert str(validation.sqref) == "I2"


def test_meta_values_round_trip(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"
    metadata = valid_metadata()

    export_call_list(output_path, [], metadata)

    meta_sheet = load_workbook(output_path)["_meta"]

    exported = {
        meta_sheet.cell(row=row_number, column=1).value: meta_sheet.cell(
            row=row_number, column=2
        ).value
        for row_number in range(2, meta_sheet.max_row + 1)
    }

    assert exported == metadata


def test_missing_metadata_is_rejected(tmp_path):
    metadata = valid_metadata()
    metadata.pop("week")

    with pytest.raises(ValueError, match="week"):
        export_call_list(tmp_path / "ringlista.xlsx", [], metadata)


def test_column_widths_match_contract():
    assert set(COLUMN_WIDTHS) >= set(COLUMNS)


def test_hidden_columns_follow_contract(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    export_call_list(output_path, [], valid_metadata())

    worksheet = load_workbook(output_path)[SHEET]

    for column_name in HIDDEN:
        column_letter = worksheet.cell(row=1, column=COL[column_name]).column_letter
        assert worksheet.column_dimensions[column_letter].hidden is True


def test_export_round_trip_preserves_rows_and_metadata(tmp_path):
    output_path = tmp_path / "ringlista.xlsx"

    rows = [
        CallListRow(
            row_id=17,
            rank=1,
            score=2.0,
            salon="Klipp & Co",
            area="Stockholm",
            phone="+46701234567",
            source="https://example.com/contact",
            reasons="Nystartad salong",
        ),
        CallListRow(
            row_id=18,
            rank=2,
            score=1.0,
            salon="Studio Sax",
            area="Uppsala",
            phone="+46707654321",
            source="https://example.com",
            reasons="1-4 anställda",
        ),
    ]

    metadata = valid_metadata()

    export_call_list(output_path, rows, metadata)

    workbook = load_workbook(output_path)
    worksheet = workbook[SHEET]

    exported_row_ids = [
        worksheet.cell(row=row_number, column=COL["row_id"]).value
        for row_number in range(2, 2 + len(rows))
    ]

    assert exported_row_ids == [17, 18]

    meta_sheet = workbook[META_SHEET]
    exported_metadata = {
        meta_sheet.cell(row=row_number, column=1).value: meta_sheet.cell(
            row=row_number, column=2
        ).value
        for row_number in range(2, meta_sheet.max_row + 1)
    }

    assert exported_metadata == metadata
