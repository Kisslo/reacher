import importlib

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from reacher.cli import FIXTURES_DIR, app
from reacher.db import connect, migrate
from reacher.excel.contract import COL, SHEET
from reacher.excel.export import CallListRow, export_call_list
from reacher.simulation import SimulationError, simulate_outcomes

outcome_import = importlib.import_module("reacher.excel.import")


@pytest.fixture
def prepared(tmp_path):
    db_path = tmp_path / "simulation.db"
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "simulated.xlsx"
    conn = connect(db_path)
    migrate(conn)

    salon_ids = []
    for orgnr, cfar, name in [
        ("0107053316", "40000005", "Fixture 1"),
        ("5560002023", "40001234", "Fixture 2"),
        ("5560003039", None, "Fixture 3"),
    ]:
        salon_ids.append(
            conn.execute(
                "INSERT INTO salon (orgnr, cfar, name, first_seen_at, last_seen_at) "
                "VALUES (?, ?, ?, '2026-09-29', '2026-09-29')",
                (orgnr, cfar, name),
            ).lastrowid
        )

    call_list_id = conn.execute(
        "INSERT INTO call_list (week, seller, scoring_version, created_at) "
        "VALUES ('2026w40', 'Anna', 'test-v1', '2026-09-29')"
    ).lastrowid
    row_ids = []
    for rank, salon_id in enumerate(salon_ids, start=1):
        row_ids.append(
            conn.execute(
                "INSERT INTO call_list_row "
                "(call_list_id, salon_id, rank, score, reasons, phone) "
                "VALUES (?, ?, ?, 1, '[]', ?)",
                (call_list_id, salon_id, rank, f"+461234567{rank}"),
            ).lastrowid
        )
    conn.commit()

    export_call_list(
        input_path,
        [
            CallListRow(
                row_id,
                rank,
                1,
                f"Fixture {rank}",
                "Stockholm",
                f"+461234567{rank}",
                "",
                "",
            )
            for rank, row_id in enumerate(row_ids, start=1)
        ],
        {
            "call_list_id": str(call_list_id),
            "week": "2026w40",
            "seller": "Anna",
            "scoring_version": "test-v1",
            "generated_at": "2026-09-29T10:00:00+00:00",
        },
    )
    yield conn, db_path, input_path, output_path, call_list_id
    conn.close()


def outcomes(path):
    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = workbook[SHEET]
        return [
            worksheet.cell(row=row, column=COL["Utfall"]).value
            for row in range(2, worksheet.max_row + 1)
        ]
    finally:
        workbook.close()


def test_seed_is_reproducible_and_guarantees_demo_cases(prepared):
    conn, _, input_path, output_path, _ = prepared
    other_path = output_path.with_name("other.xlsx")

    first = simulate_outcomes(conn, input_path, output_path, FIXTURES_DIR / "ground_truth.csv", 7)
    second = simulate_outcomes(conn, input_path, other_path, FIXTURES_DIR / "ground_truth.csv", 7)

    assert first == second
    assert outcomes(output_path) == outcomes(other_path)
    assert "Ej nådd" in outcomes(output_path)
    assert "Spärra" in outcomes(output_path)


def test_all_unreachable_seed_still_gets_a_block(prepared):
    conn, _, input_path, output_path, _ = prepared

    simulate_outcomes(conn, input_path, output_path, FIXTURES_DIR / "ground_truth.csv", 149)

    generated = outcomes(output_path)
    assert "Ej nådd" in generated
    assert "Spärra" in generated


def test_existing_output_is_rejected_without_overwriting(prepared):
    conn, _, input_path, output_path, _ = prepared
    original = b"existing workbook placeholder"
    output_path.write_bytes(original)

    with pytest.raises(SimulationError, match="Output path already exists"):
        simulate_outcomes(conn, input_path, output_path, FIXTURES_DIR / "ground_truth.csv", 7)

    assert output_path.read_bytes() == original


def test_unknown_salon_is_rejected_without_output(prepared):
    conn, _, input_path, output_path, call_list_id = prepared
    salon_id = conn.execute(
        "INSERT INTO salon (orgnr, cfar, name, first_seen_at, last_seen_at) "
        "VALUES ('9999999999', '99999999', 'Real salon', '2026-09-29', '2026-09-29')"
    ).lastrowid
    row_id = conn.execute(
        "INSERT INTO call_list_row "
        "(call_list_id, salon_id, rank, score, reasons, phone) "
        "VALUES (?, ?, 4, 1, '[]', '+4612345678')",
        (call_list_id, salon_id),
    ).lastrowid
    conn.commit()
    workbook = load_workbook(input_path)
    workbook[SHEET].cell(row=5, column=COL["row_id"]).value = row_id
    workbook.save(input_path)

    with pytest.raises(SimulationError, match="not in the fixtures"):
        simulate_outcomes(conn, input_path, output_path, FIXTURES_DIR / "ground_truth.csv", 7)
    assert not output_path.exists()


def test_generated_file_imports_through_t2_04(prepared):
    conn, _, input_path, output_path, _ = prepared
    simulate_outcomes(conn, input_path, output_path, FIXTURES_DIR / "ground_truth.csv", 11)

    summary = outcome_import.import_outcomes(conn, output_path)

    assert summary.imported == 3
    assert summary.unknown_outcomes == ()
    assert summary.unknown_row_ids == ()


def test_cli_simulates_to_requested_output(prepared):
    conn, db_path, input_path, output_path, _ = prepared
    conn.close()

    result = CliRunner().invoke(
        app,
        [
            "simulate-outcomes",
            str(input_path),
            str(output_path),
            "--seed",
            "5",
            "--db",
            str(db_path),
        ],
    )

    assert result.exit_code == 0, result.stdout
    assert output_path.exists()
    assert "seed 5" in result.stdout
