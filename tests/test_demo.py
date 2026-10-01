"""J-03: the Friday demo from README.md, run end to end through the CLI."""

import shutil
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from reacher.cli import app
from reacher.db import connect
from reacher.excel.contract import COL, FIRST_DATA_ROW, SHEET, Outcome

REPO_ROOT = Path(__file__).resolve().parents[1]
WEEK = "2026w40"
NEXT_WEEK = "2026w41"
SELLERS = ("Anna", "Bengt")

# Must stay identical to the "Fredagsdemo" section in README.md.
DEMO_STEPS = [
    ["init-db"],
    ["load-seed"],
    ["build-lists", WEEK, "Anna,Bengt"],
    *(
        ["simulate-outcomes", f"output/{WEEK}_{seller}.xlsx", f"output/{WEEK}_{seller}_utfall.xlsx"]
        for seller in SELLERS
    ),
    *(["import-outcomes", f"output/{WEEK}_{seller}_utfall.xlsx"] for seller in SELLERS),
    ["build-lists", NEXT_WEEK, "Anna,Bengt"],
    ["report", WEEK],
]


def listed_salons(conn, week):
    """(salon_id, orgnr) for every row in the xlsx files sent out for ``week``."""
    salons = set()
    for seller in SELLERS:
        workbook = load_workbook(Path("output") / f"{week}_{seller}.xlsx", read_only=True)
        try:
            row_ids = [
                values[COL["row_id"] - 1]
                for values in workbook[SHEET].iter_rows(min_row=FIRST_DATA_ROW, values_only=True)
            ]
        finally:
            workbook.close()
        for row_id in row_ids:
            row = conn.execute(
                "SELECT salon.id, salon.orgnr FROM call_list_row "
                "JOIN salon ON salon.id = call_list_row.salon_id "
                "JOIN call_list ON call_list.id = call_list_row.call_list_id "
                "WHERE call_list_row.id = ? AND call_list.week = ?",
                (row_id, week),
            ).fetchone()
            assert row is not None, f"row_id {row_id} in the {week} file is not in that week's list"
            salons.add((row["id"], row["orgnr"]))
    return salons


def orgnrs_with_outcome(conn, week, outcome):
    return {
        row["orgnr"]
        for row in conn.execute(
            "SELECT DISTINCT salon.orgnr FROM outcome "
            "JOIN call_list_row ON call_list_row.id = outcome.call_list_row_id "
            "JOIN call_list ON call_list.id = call_list_row.call_list_id "
            "JOIN salon ON salon.id = call_list_row.salon_id "
            "WHERE call_list.week = ? AND outcome.outcome = ?",
            (week, outcome.value),
        )
    }


def test_friday_demo_runs_the_full_loop(tmp_path, monkeypatch):
    # build-lists reads scoring.yaml and writes output/ in the working directory.
    shutil.copy(REPO_ROOT / "scoring.yaml", tmp_path)
    monkeypatch.chdir(tmp_path)

    runner = CliRunner()
    for step in DEMO_STEPS:
        result = runner.invoke(app, step)
        assert result.exit_code == 0, f"reacher {' '.join(step)}:\n{result.output}"
    report_output = result.output

    with closing(connect(tmp_path / "reacher.db")) as conn:
        blocked = orgnrs_with_outcome(conn, WEEK, Outcome.SPARRA)
        registered = orgnrs_with_outcome(conn, WEEK, Outcome.REGISTRERAD)
        this_week = listed_salons(conn, WEEK)
        next_week = listed_salons(conn, NEXT_WEEK)

    assert blocked, "the simulation guarantees at least one Spärra per list"
    assert not {orgnr for _, orgnr in next_week} & blocked
    assert not {orgnr for _, orgnr in next_week} & registered
    # Nothing else changed between the weeks, so only the suppressed companies drop out.
    assert next_week == {
        (salon_id, orgnr) for salon_id, orgnr in this_week if orgnr not in blocked | registered
    }

    assert f"Week {WEEK}" in report_output
    assert "All lists:" in report_output
    assert "sent 20" in report_output
