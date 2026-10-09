import json
from contextlib import closing
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from reacher.cli import FIXTURES_DIR
from reacher.cli import build_lists as cli_build_lists
from reacher.config import ScoringConfig
from reacher.db import connect, migrate
from reacher.excel.contract import COL, FIRST_DATA_ROW, SHEET
from reacher.ingest import ingest, ingest_financials
from reacher.lists import build_call_lists
from reacher.sources.csv_source import CsvFinancialSource, CsvSource

TODAY = date(2026, 9, 28)
FINANCIAL_SIGNALS = ("loss_making", "low_revenue", "declining_revenue")


def make_config(small_employer_weight=1, financial_enabled=True) -> ScoringConfig:
    return ScoringConfig.model_validate(
        {
            "version": "test-v1",
            "half_life_days": 90,
            "signals": {
                "registered_recently": {"enabled": True, "weight": 1, "months": 24},
                "small_employer": {
                    "enabled": True,
                    "weight": small_employer_weight,
                    "classes": ["2"],
                },
                "loss_making": {"enabled": financial_enabled, "weight": 0},
                "low_revenue": {"enabled": financial_enabled, "weight": 0, "below_sek": 500000},
                "declining_revenue": {"enabled": financial_enabled, "weight": 0, "years": 3},
            },
        }
    )


CONFIG = make_config()


@pytest.fixture
def conn(tmp_path):
    with closing(connect(tmp_path / "lists.db")) as connection:
        migrate(connection)
        yield connection


def add_salon(
    conn,
    name,
    *,
    registered_at=None,
    employee_class=None,
    phone=None,
    city="Stockholm",
    street=None,
    postal_code=None,
):
    cursor = conn.execute(
        "INSERT INTO salon "
        "(orgnr, name, street, postal_code, city, registered_at, employee_class, legal_form, "
        "ftax_status, vat_status, employer_status, company_status, workplace_status, "
        "ad_block_type, phone_block_type, workplace_ad_block_type, workplace_phone_block_type, "
        "first_seen_at, last_seen_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, '49', '1', '1', '1', '1', '1', '1', '1', '1', '1', ?, ?)",
        (
            name,
            name,
            street,
            postal_code,
            city,
            registered_at,
            employee_class,
            "2026-09-28",
            "2026-09-28",
        ),
    )
    salon_id = cursor.lastrowid
    if phone:
        conn.execute(
            "INSERT INTO contact (salon_id, kind, value, source_url, found_at) "
            "VALUES (?, 'phone', ?, ?, '2026-09-28')",
            (salon_id, phone, f"https://example.test/{name}"),
        )
    conn.commit()
    return salon_id


def test_build_lists_ranks_splits_and_counts_missing_phones(conn):
    add_salon(
        conn,
        "New Small",
        registered_at="2026-01-01",
        employee_class="2",
        phone="+461",
    )
    add_salon(conn, "New", registered_at="2026-01-01", phone="+462")
    add_salon(conn, "Small", employee_class="2", phone="+463")
    add_salon(conn, "Neither", phone="+464")
    add_salon(conn, "No Phone", registered_at="2026-01-01", employee_class="2")

    built = build_call_lists(conn, "2026w40", ("anna", "bengt"), CONFIG, TODAY)

    assert [[row.rank for row in item.rows] for item in built] == [[1, 3], [2, 4]]
    assert [[row.salon for row in item.rows] for item in built] == [
        ["New Small", "Small"],
        ["New", "Neither"],
    ]
    assert all(item.skipped_without_phone == 1 for item in built)


def test_snapshot_freezes_score_reasons_signals_and_phone(conn):
    salon_id = add_salon(
        conn,
        "Frozen Salon",
        registered_at="2026-01-01",
        employee_class="2",
        phone="+461234",
        city="Stockholm",
    )

    built = build_call_lists(conn, "2026w40", ("anna",), CONFIG, TODAY)
    original = built[0].rows[0]

    conn.execute(
        "UPDATE salon SET city = 'Gothenburg', registered_at = '2010-01-01', "
        "employee_class = '1' WHERE id = ?",
        (salon_id,),
    )
    conn.execute("UPDATE contact SET value = '+469999' WHERE salon_id = ?", (salon_id,))
    stored = conn.execute(
        "SELECT score, reasons, signals, phone FROM call_list_row WHERE id = ?",
        (original.row_id,),
    ).fetchone()

    assert stored["score"] == 2
    assert json.loads(stored["reasons"]) == [
        "Registrerad för 8 månader sedan",
        "1-4 anställda",
    ]
    assert json.loads(stored["signals"]) == ["registered_recently", "small_employer"]
    assert stored["phone"] == "+461234"
    assert original.town == "Stockholm"


def test_snapshot_includes_workplace_address_and_financial_years(conn):
    salon_id = add_salon(
        conn,
        "Financial Salon",
        phone="+461234",
        street="Besöksgatan 4",
        postal_code="11455",
    )
    conn.execute(
        "INSERT INTO financial_fact "
        "(orgnr, period_end, key, value, source_document, fetched_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        ("5590001010", "2025-04-30", "revenue", 965400, "doc-1", "2026-09-28"),
    )
    conn.execute(
        "INSERT INTO financial_fact "
        "(orgnr, period_end, key, value, source_document, fetched_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        ("5590001010", "2025-04-30", "net_result", -12600, "doc-1", "2026-09-28"),
    )
    conn.execute("UPDATE salon SET orgnr = ? WHERE id = ?", ("5590001010", salon_id))
    conn.commit()

    built = build_call_lists(conn, "2026w40", ("anna",), CONFIG, TODAY)
    row = built[0].rows[0]

    assert row.address == "Besöksgatan 4, 11455"
    assert row.town == "Stockholm"
    assert row.revenue == "2024/25: 965 400 kr"
    assert row.result == "2024/25: -12 600 kr"

    stored = conn.execute(
        "SELECT address, town, revenue, result FROM call_list_row WHERE id = ?",
        (row.row_id,),
    ).fetchone()
    assert tuple(stored) == (
        "Besöksgatan 4, 11455",
        "Stockholm",
        "2024/25: 965 400 kr",
        "2024/25: -12 600 kr",
    )


def test_salon_without_signals_stores_an_empty_list_not_null(conn):
    """NULL betyder "byggd före T2-08, okänt". En rad utan signaler är []."""
    add_salon(conn, "Plain Salon", phone="+461")

    built = build_call_lists(conn, "2026w40", ("anna",), CONFIG, TODAY)

    stored = conn.execute(
        "SELECT signals FROM call_list_row WHERE id = ?", (built[0].rows[0].row_id,)
    ).fetchone()
    assert json.loads(stored["signals"]) == []


def test_shadow_signal_is_stored_but_not_shown(conn, tmp_path, monkeypatch):
    """weight: 0 (D26): signalen fryses i call_list_row.signals så att T2-07 kan
    mäta den, men den ger inga poäng och syns inte i "Varför vi ringer"."""
    add_salon(conn, "Shadow Salon", employee_class="2", phone="+461")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        "reacher.config.ScoringConfig.load", lambda: make_config(small_employer_weight=0)
    )

    cli_build_lists("2026w40", "anna", tmp_path / "lists.db")

    stored = conn.execute("SELECT score, reasons, signals FROM call_list_row").fetchone()
    assert stored["score"] == 0
    assert json.loads(stored["reasons"]) == []
    assert json.loads(stored["signals"]) == ["small_employer"]

    sheet = load_workbook(tmp_path / "output" / "2026w40_anna.xlsx")[SHEET]
    assert not sheet.cell(row=FIRST_DATA_ROW, column=COL["Varför vi ringer"]).value
    assert sheet.cell(row=FIRST_DATA_ROW, column=COL["Salong"]).value == "Shadow Salon"


def load_fixtures(conn):
    ingest(conn, CsvSource(FIXTURES_DIR).fetch(), now="2026-09-28T08:00:00+00:00")
    ingest_financials(
        conn, CsvFinancialSource(FIXTURES_DIR).fetch(), now="2026-09-28T08:00:00+00:00"
    )


def frozen_signals(conn, call_list_ids) -> dict[str, set[str]]:
    """Signalnyckel -> namnen på salongerna som fick den, ur call_list_row."""
    placeholders = ",".join("?" * len(call_list_ids))
    by_signal: dict[str, set[str]] = {}
    for row in conn.execute(
        "SELECT salon.name, call_list_row.signals FROM call_list_row "
        "JOIN salon ON salon.id = call_list_row.salon_id "
        f"WHERE call_list_row.call_list_id IN ({placeholders})",
        call_list_ids,
    ):
        for key in json.loads(row["signals"]):
            by_signal.setdefault(key, set()).add(row["name"])
    return by_signal


def test_financial_shadow_signals_are_frozen_per_row_from_the_fixtures(conn):
    """Läses via latest_financial_fact på orgnr. Se tests/fixtures/README.md."""
    load_fixtures(conn)

    built = build_call_lists(conn, "2026w40", ("anna", "bengt"), CONFIG, TODAY)

    signals = frozen_signals(conn, [item.call_list_id for item in built])
    # Brynboden: net_result saknas senaste året. Frisyr & Form: förlusten är ett
    # äldre år. Nackeateljén: resultat exakt 0. Saxateljén: exakt 500 000.
    assert signals["loss_making"] == {"Salong Dubbelgångaren", "Franssalongen"}
    assert signals["low_revenue"] == {"Franssalongen", "Nackeateljén"}
    # Hårateljén saknar 2023, Nackeateljén och Brynboden har bara två år.
    assert signals["declining_revenue"] == {"Salong Dubbelgångaren"}


def test_financial_shadow_signals_do_not_change_the_lists(conn):
    """Skuggläge (D26): samma rangordning, poäng och skäl som utan signalerna."""
    load_fixtures(conn)

    def lists(week, config):
        return [
            [(row.salon_id, row.rank, row.score, row.reasons) for row in item.rows]
            for item in build_call_lists(conn, week, ("anna", "bengt"), config, TODAY)
        ]

    assert lists("2026w40", CONFIG) == lists("2026w41", make_config(financial_enabled=False))


def test_duplicate_week_and_seller_rolls_back(conn):
    add_salon(conn, "Salon", phone="+461")
    build_call_lists(conn, "2026w40", ("anna",), CONFIG, TODAY)

    with pytest.raises(Exception, match="UNIQUE"):
        build_call_lists(conn, "2026w40", ("anna",), CONFIG, TODAY)

    assert conn.execute("SELECT count(*) FROM call_list").fetchone()[0] == 1


def test_sellers_must_be_nonempty_and_unique(conn):
    with pytest.raises(ValueError, match="At least one"):
        build_call_lists(conn, "2026w40", (), CONFIG, TODAY)
    with pytest.raises(ValueError, match="unique"):
        build_call_lists(conn, "2026w40", ("anna", "anna"), CONFIG, TODAY)


def test_cli_exports_workbook_and_persists_path(conn, tmp_path, monkeypatch):
    add_salon(conn, "Exported Salon", phone="+461")
    db_path = tmp_path / "lists.db"
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("reacher.config.ScoringConfig.load", lambda: CONFIG)

    cli_build_lists("2026w40", "anna", db_path)

    output_path = tmp_path / "output" / "2026w40_anna.xlsx"
    assert output_path.exists()
    stored_path = conn.execute("SELECT file_path FROM call_list").fetchone()[0]
    assert Path(stored_path) == Path("output") / "2026w40_anna.xlsx"

    workbook = load_workbook(output_path)
    metadata = {
        workbook["_meta"].cell(row=row_number, column=1).value: workbook["_meta"]
        .cell(row=row_number, column=2)
        .value
        for row_number in range(2, workbook["_meta"].max_row + 1)
    }
    assert metadata["seller"] == "anna"
    assert metadata["week"] == "2026w40"
    assert metadata["scoring_version"] == "test-v1"
