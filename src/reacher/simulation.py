"""Generate fixture-only outcomes for development demos (T2-05)."""

import csv
import random
import sqlite3
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from reacher.excel.contract import (
    COL,
    FIRST_DATA_ROW,
    HEADER_ROW,
    META_KEYS,
    META_SHEET,
    SHEET,
    Outcome,
)

OUTCOME_WEIGHTS = {
    True: (
        (Outcome.EJ_NADD.value, 20),
        (Outcome.INTRESSERAD.value, 40),
        (Outcome.REGISTRERAD.value, 20),
        (Outcome.NEJ.value, 15),
        (Outcome.SPARRA.value, 5),
    ),
    False: (
        (Outcome.EJ_NADD.value, 30),
        (Outcome.INTRESSERAD.value, 8),
        (Outcome.REGISTRERAD.value, 4),
        (Outcome.NEJ.value, 48),
        (Outcome.SPARRA.value, 10),
    ),
}

COMMENTS = {
    Outcome.EJ_NADD.value: "Simulerat: inget svar",
    Outcome.INTRESSERAD.value: "Simulerat: kunden var intresserad",
    Outcome.REGISTRERAD.value: "Simulerat: kunden registrerades",
    Outcome.NEJ.value: "Simulerat: kunden var inte intresserad",
    Outcome.SPARRA.value: "Simulerat: kunden bad att inte bli kontaktad",
}


class SimulationError(ValueError):
    """The workbook is not a safe fixture workbook to simulate."""


def _fixture_key(orgnr: str, cfar: str | None) -> tuple[str, str]:
    return orgnr, cfar or ""


def _read_ground_truth(path: Path) -> dict[tuple[str, str], bool]:
    with path.open(encoding="utf-8", newline="") as stream:
        rows = csv.DictReader(stream)
        expected = ("orgnr", "cfar", "has_empty_chairs")
        if tuple(rows.fieldnames or ()) != expected:
            raise SimulationError("ground_truth.csv headers do not match the fixture contract")

        truth: dict[tuple[str, str], bool] = {}
        for row in rows:
            value = row["has_empty_chairs"]
            if value not in {"0", "1"}:
                raise SimulationError("ground_truth.csv contains a non-boolean value")
            key = _fixture_key(row["orgnr"], row["cfar"] or None)
            if key in truth:
                raise SimulationError(f"ground_truth.csv contains duplicate salon {key}")
            truth[key] = value == "1"
        return truth


def _metadata(workbook) -> dict[str, str]:
    if META_SHEET not in workbook.sheetnames:
        raise SimulationError(f"Workbook is missing {META_SHEET!r} sheet")

    sheet = workbook[META_SHEET]
    values = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value is not None
    }
    missing = [key for key in META_KEYS if key not in values]
    if missing:
        raise SimulationError(f"Workbook is missing metadata keys: {missing}")
    return {key: str(values[key]) for key in META_KEYS}


def _row_id(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    try:
        converted = int(value)
    except (TypeError, ValueError):
        return None
    return converted if converted == value or str(converted) == str(value).strip() else None


def _call_list_id(metadata: dict[str, str]) -> int:
    try:
        return int(metadata["call_list_id"])
    except (TypeError, ValueError) as error:
        raise SimulationError("Metadata call_list_id must be an integer") from error


def _validate_headers(worksheet) -> None:
    for column_name, column_number in COL.items():
        if worksheet.cell(row=HEADER_ROW, column=column_number).value != column_name:
            raise SimulationError("Workbook headers do not match the Excel contract")


def _row_values(worksheet) -> Iterable[tuple[int, int]]:
    for excel_row in range(FIRST_DATA_ROW, worksheet.max_row + 1):
        row_id = _row_id(worksheet.cell(row=excel_row, column=COL["row_id"]).value)
        if row_id is None:
            raise SimulationError(f"Workbook row {excel_row} has an invalid row_id")
        yield excel_row, row_id


def _choose_outcome(rng: random.Random, has_empty_chairs: bool) -> str:
    outcomes, weights = zip(*OUTCOME_WEIGHTS[has_empty_chairs], strict=True)
    return rng.choices(outcomes, weights=weights, k=1)[0]


def simulate_outcomes(
    conn: sqlite3.Connection,
    input_path: Path,
    output_path: Path,
    ground_truth_path: Path,
    seed: int,
) -> dict[str, int]:
    """Fill a copy of a fixture call-list workbook with seeded outcomes.

    Every row is validated against ``ground_truth.csv`` before the output is
    written, so a real or partially unknown call list cannot be simulated.
    """
    if input_path.resolve() == output_path.resolve():
        raise SimulationError("Input and output paths must be different")
    if output_path.exists():
        raise SimulationError(f"Output path already exists: {output_path}")

    truth = _read_ground_truth(ground_truth_path)
    workbook = load_workbook(input_path)
    try:
        if SHEET not in workbook.sheetnames:
            raise SimulationError(f"Workbook is missing {SHEET!r} sheet")

        metadata = _metadata(workbook)
        call_list_id = _call_list_id(metadata)
        call_list = conn.execute(
            "SELECT week FROM call_list WHERE id = ?", (call_list_id,)
        ).fetchone()
        if call_list is None or call_list["week"] != metadata["week"]:
            raise SimulationError(
                f"Metadata does not match call list {call_list_id} and week {metadata['week']!r}"
            )

        worksheet = workbook[SHEET]
        _validate_headers(worksheet)
        rows: list[tuple[int, bool]] = []
        seen_row_ids: set[int] = set()
        for excel_row, row_id in _row_values(worksheet):
            if row_id in seen_row_ids:
                raise SimulationError(f"Workbook contains duplicate row_id {row_id}")
            seen_row_ids.add(row_id)
            salon = conn.execute(
                "SELECT salon.orgnr, salon.cfar FROM call_list_row "
                "JOIN salon ON salon.id = call_list_row.salon_id "
                "WHERE call_list_row.id = ? AND call_list_row.call_list_id = ?",
                (row_id, call_list_id),
            ).fetchone()
            if salon is None:
                raise SimulationError(f"row_id {row_id} is not in call list {call_list_id}")
            key = _fixture_key(salon["orgnr"], salon["cfar"])
            if key not in truth:
                raise SimulationError(f"Salon {key} is not in the fixtures")
            rows.append((excel_row, truth[key]))

        if len(rows) < 2:
            raise SimulationError("At least two fixture salons are required for simulation")

        rng = random.Random(seed)
        generated = [
            [_choose_outcome(rng, has_empty_chairs), excel_row]
            for excel_row, has_empty_chairs in rows
        ]

        ej_nadd_indices = [
            index for index, row in enumerate(generated) if row[0] == Outcome.EJ_NADD.value
        ]
        if not ej_nadd_indices:
            protected_ej_index = rng.randrange(len(generated))
            generated[protected_ej_index][0] = Outcome.EJ_NADD.value
        else:
            protected_ej_index = ej_nadd_indices[0]
        if not any(row[0] == Outcome.SPARRA.value for row in generated):
            candidates = [index for index in range(len(generated)) if index != protected_ej_index]
            index = rng.choice(candidates)
            generated[index][0] = Outcome.SPARRA.value

        counts: dict[str, int] = {}
        for outcome, excel_row in generated:
            worksheet.cell(row=excel_row, column=COL["Utfall"]).value = outcome
            worksheet.cell(row=excel_row, column=COL["Kommentar"]).value = COMMENTS[outcome]
            counts[outcome] = counts.get(outcome, 0) + 1

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return counts
    finally:
        workbook.close()
