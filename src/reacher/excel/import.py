"""Import salesperson outcomes from a returned call-list workbook."""

import sqlite3
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from reacher.db import now
from reacher.excel.contract import (
    COL,
    FIRST_DATA_ROW,
    HEADER_ROW,
    META_KEYS,
    META_SHEET,
    SHEET,
    Outcome,
)

SUSPICIOUS_COMMENT_TERMS = ("spärra", "ring inte", "ej kontakt")


@dataclass(frozen=True)
class OutcomeImportSummary:
    """Results from one workbook import."""

    imported: int
    outcomes: dict[str, int]
    empty_outcomes: int
    unknown_outcomes: tuple[str, ...]
    unknown_row_ids: tuple[Any, ...]
    suppressions_added: dict[str, int]
    manual_review_row_ids: tuple[int, ...]


class OutcomeImportError(ValueError):
    """The workbook cannot be associated with the requested call list."""


def _metadata(workbook) -> dict[str, str]:
    if META_SHEET not in workbook.sheetnames:
        raise OutcomeImportError(f"Workbook is missing {META_SHEET!r} sheet")

    sheet = workbook[META_SHEET]
    values = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value is not None
    }
    missing = [key for key in META_KEYS if key not in values]
    if missing:
        raise OutcomeImportError(f"Workbook is missing metadata keys: {missing}")
    return {key: str(values[key]) for key in META_KEYS}


def _call_list_id(metadata: dict[str, str]) -> int:
    try:
        return int(metadata["call_list_id"])
    except (TypeError, ValueError) as error:
        raise OutcomeImportError("Metadata call_list_id must be an integer") from error


def _normalise_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _row_id(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    try:
        converted = int(value)
    except (TypeError, ValueError):
        return None
    return converted if converted == value or str(converted) == str(value).strip() else None


def import_outcomes(conn: sqlite3.Connection, workbook_path: Path) -> OutcomeImportSummary:
    """Import valid outcomes and comments from ``workbook_path``.

    Rows are matched by ``row_id`` and must belong to the call list named in
    ``_meta``. Invalid outcome values and unknown row IDs are reported and
    skipped; metadata errors reject the workbook before any write occurs.
    """
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if SHEET not in workbook.sheetnames:
            raise OutcomeImportError(f"Workbook is missing {SHEET!r} sheet")

        metadata = _metadata(workbook)
        call_list_id = _call_list_id(metadata)
        call_list = conn.execute(
            "SELECT week FROM call_list WHERE id = ?", (call_list_id,)
        ).fetchone()
        if call_list is None or call_list["week"] != metadata["week"]:
            raise OutcomeImportError(
                f"Metadata does not match call list {call_list_id} and week {metadata['week']!r}"
            )

        worksheet = workbook[SHEET]
        expected_headers = {
            column: worksheet.cell(row=HEADER_ROW, column=index).value
            for column, index in COL.items()
        }
        if any(expected_headers[column] != column for column in COL):
            raise OutcomeImportError("Workbook headers do not match the Excel contract")

        valid_outcomes = {outcome.value for outcome in Outcome}
        outcome_counts: Counter[str] = Counter()
        unknown_outcomes: set[str] = set()
        unknown_row_ids: set[Any] = set()
        manual_review_row_ids: set[int] = set()
        rows: list[tuple[int, str, str, str]] = []

        for values in worksheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            row_id = _row_id(values[COL["row_id"] - 1])
            outcome = _normalise_cell(values[COL["Utfall"] - 1])
            comment = _normalise_cell(values[COL["Kommentar"] - 1])

            if comment and outcome != Outcome.SPARRA.value:
                lowered = comment.casefold()
                if any(term in lowered for term in SUSPICIOUS_COMMENT_TERMS):
                    if row_id is not None:
                        manual_review_row_ids.add(row_id)

            if not outcome:
                continue
            if outcome not in valid_outcomes:
                unknown_outcomes.add(outcome)
                continue
            if row_id is None:
                unknown_row_ids.add(values[COL["row_id"] - 1])
                continue

            row = conn.execute(
                "SELECT salon.orgnr FROM call_list_row "
                "JOIN salon ON salon.id = call_list_row.salon_id "
                "WHERE call_list_row.id = ? AND call_list_row.call_list_id = ?",
                (row_id, call_list_id),
            ).fetchone()
            if row is None:
                unknown_row_ids.add(row_id)
                continue

            rows.append((row_id, outcome, comment, row["orgnr"]))
            outcome_counts[outcome] += 1

        suppressions_added: Counter[str] = Counter()
        with conn:
            imported_at = now()
            for row_id, outcome, comment, orgnr in rows:
                conn.execute(
                    "INSERT INTO outcome (call_list_row_id, outcome, comment, imported_at) "
                    "VALUES (?, ?, ?, ?) "
                    "ON CONFLICT(call_list_row_id) DO UPDATE SET "
                    "outcome = excluded.outcome, comment = excluded.comment, "
                    "imported_at = excluded.imported_at",
                    (row_id, outcome, comment or None, imported_at),
                )

                reason = {
                    Outcome.SPARRA.value: "opt_out",
                    Outcome.REGISTRERAD.value: "existing_customer",
                }.get(outcome)
                if reason is None:
                    continue
                cursor = conn.execute(
                    "INSERT INTO suppression (orgnr, reason, created_at) VALUES (?, ?, ?) "
                    "ON CONFLICT(orgnr, reason) DO NOTHING",
                    (orgnr, reason, imported_at),
                )
                suppressions_added[reason] += cursor.rowcount

        return OutcomeImportSummary(
            imported=len(rows),
            outcomes=dict(outcome_counts),
            empty_outcomes=sum(
                1
                for values in worksheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True)
                if not _normalise_cell(values[COL["Utfall"] - 1])
            ),
            unknown_outcomes=tuple(sorted(unknown_outcomes)),
            unknown_row_ids=tuple(sorted(unknown_row_ids, key=str)),
            suppressions_added=dict(suppressions_added),
            manual_review_row_ids=tuple(sorted(manual_review_row_ids)),
        )
    finally:
        workbook.close()
