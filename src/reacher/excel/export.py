from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from reacher.excel.contract import (
    COL,
    COLUMN_WIDTHS,
    COLUMNS,
    EDITABLE,
    FIRST_DATA_ROW,
    HEADER_ROW,
    HIDDEN,
    META_KEYS,
    META_SHEET,
    SHEET,
    Outcome,
)


@dataclass(frozen=True)
class CallListRow:
    row_id: int
    rank: int
    score: float
    salon: str
    area: str
    phone: str
    source: str
    reasons: str


def export_call_list(
    output_path: Path,
    rows: Sequence[CallListRow],
    metadata: Mapping[str, str],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET

    for column_number, column_name in enumerate(COLUMNS, start=1):
        worksheet.cell(
            row=HEADER_ROW,
            column=column_number,
            value=column_name,
        )

        worksheet.column_dimensions[get_column_letter(column_number)].width = COLUMN_WIDTHS.get(
            column_name, 18
        )

    for column_name in HIDDEN:
        column_number = COL[column_name]
        worksheet.column_dimensions[get_column_letter(column_number)].hidden = True

    for excel_row, row in enumerate(rows, start=FIRST_DATA_ROW):
        values = {
            "row_id": row.row_id,
            "Rang": row.rank,
            "Poäng": row.score,
            "Salong": row.salon,
            "Område": row.area,
            "Telefon": row.phone,
            "Källa": row.source,
            "Varför vi ringer": row.reasons,
            "Utfall": None,
            "Kommentar": None,
        }

        for column_number, column_name in enumerate(COLUMNS, start=1):
            value = values.get(column_name)
            cell = worksheet.cell(
                row=excel_row,
                column=column_number,
                value=value,
            )
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            if column_name == "Källa" and row.source:
                cell.hyperlink = row.source
                cell.style = "Hyperlink"

    outcome_values = ",".join(outcome.value for outcome in Outcome)

    validation = DataValidation(
        type="list",
        formula1=f'"{outcome_values}"',
        allow_blank=True,
    )

    worksheet.add_data_validation(validation)

    last_row = FIRST_DATA_ROW + len(rows) - 1

    if rows:
        outcome_column = COL["Utfall"]
        outcome_range = worksheet.cell(
            row=FIRST_DATA_ROW,
            column=outcome_column,
        ).coordinate

        outcome_end = worksheet.cell(
            row=last_row,
            column=outcome_column,
        ).coordinate

        validation.add(f"{outcome_range}:{outcome_end}")

    for row_number in range(FIRST_DATA_ROW, last_row + 1):
        for column_name in EDITABLE:
            cell = worksheet.cell(
                row=row_number,
                column=COL[column_name],
            )
            cell.protection = Protection(locked=False)

    worksheet.protection.sheet = True

    missing_keys = set(META_KEYS) - metadata.keys()
    if missing_keys:
        raise ValueError(f"Missing metadata keys: {sorted(missing_keys)}")

    meta_sheet = workbook.create_sheet(META_SHEET)
    meta_sheet.sheet_state = "hidden"

    meta_sheet.append(["key", "value"])

    for key in META_KEYS:
        meta_sheet.append([key, metadata[key]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
